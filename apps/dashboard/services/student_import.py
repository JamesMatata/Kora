"""Parse and upsert student roster spreadsheets without aborting the batch."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field

from django.core.exceptions import ValidationError
from django.db import transaction

from academics.models import Student

REQUIRED_COLUMNS = (
    'admission_number',
    'first_name',
    'last_name',
    'parent_name',
    'parent_phone',
)
OPTIONAL_COLUMNS = ('is_active',)
E164_RE = re.compile(r'^\+[1-9]\d{1,14}$')


@dataclass
class ImportRowResult:
    row_number: int
    status: str  # created | updated | skipped
    admission_number: str = ''
    message: str = ''


@dataclass
class ImportReport:
    created: int = 0
    updated: int = 0
    skipped: int = 0
    rows: list[ImportRowResult] = field(default_factory=list)

    @property
    def failed_rows(self):
        return [row for row in self.rows if row.status == 'skipped']


def _normalize_header(value: str) -> str:
    return (
        (value or '')
        .strip()
        .lower()
        .replace(' ', '_')
        .replace('-', '_')
    )


def _truthy_active(raw: str | None) -> bool:
    if raw is None or str(raw).strip() == '':
        return True
    value = str(raw).strip().lower()
    if value in {'1', 'true', 'yes', 'y', 'active'}:
        return True
    if value in {'0', 'false', 'no', 'n', 'inactive'}:
        return False
    raise ValueError('is_active must be yes/no, true/false, 1/0, or active/inactive.')


def _read_csv_rows(file_obj) -> list[dict]:
    raw = file_obj.read()
    if isinstance(raw, bytes):
        text = raw.decode('utf-8-sig')
    else:
        text = raw
    reader = csv.DictReader(io.StringIO(text))
    if not reader.fieldnames:
        raise ValidationError('Spreadsheet has no header row.')
    normalized = [_normalize_header(h) for h in reader.fieldnames]
    missing = [col for col in REQUIRED_COLUMNS if col not in normalized]
    if missing:
        raise ValidationError(
            f'Missing required column(s): {", ".join(missing)}.'
        )
    rows = []
    for raw_row in reader:
        mapped = {
            _normalize_header(key): (value if value is None else str(value).strip())
            for key, value in raw_row.items()
            if key is not None
        }
        rows.append(mapped)
    return rows


def _read_xlsx_rows(file_obj) -> list[dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValidationError(
            'Excel support is unavailable. Install openpyxl or upload a CSV file.'
        ) from exc

    workbook = load_workbook(file_obj, read_only=True, data_only=True)
    sheet = workbook.active
    iterator = sheet.iter_rows(values_only=True)
    try:
        header_row = next(iterator)
    except StopIteration as exc:
        raise ValidationError('Spreadsheet is empty.') from exc

    headers = [
        _normalize_header(str(cell) if cell is not None else '')
        for cell in header_row
    ]
    if not any(headers):
        raise ValidationError('Spreadsheet has no header row.')
    missing = [col for col in REQUIRED_COLUMNS if col not in headers]
    if missing:
        raise ValidationError(
            f'Missing required column(s): {", ".join(missing)}.'
        )

    rows = []
    for values in iterator:
        if values is None or all(
            cell is None or str(cell).strip() == '' for cell in values
        ):
            continue
        mapped = {}
        for idx, header in enumerate(headers):
            if not header:
                continue
            cell = values[idx] if idx < len(values) else None
            mapped[header] = '' if cell is None else str(cell).strip()
        rows.append(mapped)
    return rows


def parse_student_spreadsheet(uploaded_file) -> list[dict]:
    name = (uploaded_file.name or '').lower()
    uploaded_file.seek(0)
    if name.endswith('.xlsx'):
        return _read_xlsx_rows(uploaded_file)
    if name.endswith('.csv'):
        return _read_csv_rows(uploaded_file)
    raise ValidationError('Only .csv or .xlsx files are supported.')


def _validate_row(data: dict) -> dict:
    cleaned = {}
    missing = []
    for column in REQUIRED_COLUMNS:
        value = (data.get(column) or '').strip()
        if not value:
            missing.append(column)
        cleaned[column] = value
    if missing:
        raise ValueError(f'Missing required field(s): {", ".join(missing)}.')

    phone = cleaned['parent_phone']
    if phone.endswith('.0') and phone.replace('.', '', 1).replace('+', '').isdigit():
        raise ValueError(
            'parent_phone must be E.164 text (e.g. +254712345678). '
            'Format the spreadsheet column as text.'
        )
    if not E164_RE.match(phone):
        raise ValueError('parent_phone must be E.164 format (e.g. +254712345678).')

    cleaned['is_active'] = _truthy_active(data.get('is_active'))
    return cleaned


def import_students_for_stream(
    *,
    school,
    stream,
    rows: list[dict],
    acting_as_admin: bool,
) -> ImportReport:
    return _import_students(
        school=school,
        stream=stream,
        grade=stream.grade_level,
        rows=rows,
        acting_as_admin=acting_as_admin,
    )


def import_students_for_grade(
    *,
    school,
    grade,
    rows: list[dict],
    acting_as_admin: bool,
) -> ImportReport:
    stream = grade.ensure_default_stream()
    return _import_students(
        school=school,
        stream=stream,
        grade=grade,
        rows=rows,
        acting_as_admin=acting_as_admin,
    )


def _import_students(
    *,
    school,
    stream,
    grade,
    rows: list[dict],
    acting_as_admin: bool,
) -> ImportReport:
    report = ImportReport()

    for index, raw in enumerate(rows, start=2):  # header is row 1
        admission_preview = (raw.get('admission_number') or '').strip()
        cleaned = {}
        try:
            cleaned = _validate_row(raw)
        except ValueError as exc:
            report.skipped += 1
            report.rows.append(
                ImportRowResult(
                    row_number=index,
                    status='skipped',
                    admission_number=admission_preview,
                    message=str(exc),
                )
            )
            continue

        try:
            with transaction.atomic():
                existing = (
                    Student.objects.select_for_update()
                    .filter(
                        school=school,
                        admission_number__iexact=cleaned['admission_number'],
                    )
                    .first()
                )
                if existing is None:
                    student = Student(
                        school=school,
                        admission_number=cleaned['admission_number'],
                        first_name=cleaned['first_name'],
                        last_name=cleaned['last_name'],
                        parent_name=cleaned['parent_name'],
                        parent_phone=cleaned['parent_phone'],
                        grade_level=grade,
                        current_stream=stream,
                        is_active=cleaned['is_active'],
                    )
                    student.full_clean()
                    student.save()
                    report.created += 1
                    report.rows.append(
                        ImportRowResult(
                            row_number=index,
                            status='created',
                            admission_number=cleaned['admission_number'],
                            message='Created',
                        )
                    )
                    continue

                target_stream_id = stream.pk if stream is not None else None
                same_place = (
                    existing.current_stream_id == target_stream_id
                    and existing.grade_level_id == grade.pk
                )
                if not same_place and not acting_as_admin:
                    raise ValueError(
                        'Admission number already belongs to another class. '
                        'Ask an admin to move the student, or import as admin.'
                    )

                existing.first_name = cleaned['first_name']
                existing.last_name = cleaned['last_name']
                existing.parent_name = cleaned['parent_name']
                existing.parent_phone = cleaned['parent_phone']
                existing.is_active = cleaned['is_active']
                existing.grade_level = grade
                existing.current_stream = stream
                existing.full_clean()
                existing.save()
                report.updated += 1
                report.rows.append(
                    ImportRowResult(
                        row_number=index,
                        status='updated',
                        admission_number=cleaned['admission_number'],
                        message='Updated',
                    )
                )
        except Exception as exc:  # noqa: BLE001 — keep batch alive
            report.skipped += 1
            message = str(exc)
            if hasattr(exc, 'message_dict'):
                parts = []
                for key, errs in exc.message_dict.items():
                    parts.append(f'{key}: {", ".join(map(str, errs))}')
                message = '; '.join(parts) or message
            elif hasattr(exc, 'messages'):
                message = '; '.join(map(str, exc.messages))
            report.rows.append(
                ImportRowResult(
                    row_number=index,
                    status='skipped',
                    admission_number=cleaned.get(
                        'admission_number', admission_preview
                    ),
                    message=message,
                )
            )

    return report


TEMPLATE_HEADERS = list(REQUIRED_COLUMNS) + list(OPTIONAL_COLUMNS)
